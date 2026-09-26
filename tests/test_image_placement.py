"""Where an image actually lands: DOCX section placement, and PPTX coordinates.

Placement precision is per-format, not one property:

* PPTX is a fixed canvas - inches are inches, verified in EMU.
* DOCX flows. An image sits in a *paragraph*, so "section 3.2" is achievable and
  verifiable, but its y-position depends on everything above it, and python-docx
  has no anchored/floating image at all. A table cell is what pins a picture
  beside text.
* LaTeX PDFs float unless forced with [H] - not tested here, pdflatex is absent.
"""

from __future__ import annotations

import base64
import io
from pathlib import Path

import pytest
from PIL import Image

DOCX_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@pytest.fixture
def png(tmp_path: Path) -> Path:
    """A real, non-uniform image: a flat colour would pass a blank-frame bug."""
    buffer = io.BytesIO()
    image = Image.new("RGB", (640, 360), "white")
    for x in range(320):
        for y in range(180):
            image.putpixel((x, y), (200, 40, 40))
    image.save(buffer, format="PNG")
    path = tmp_path / "figure.png"
    path.write_bytes(buffer.getvalue())
    return path


def _long_docx(path: Path, body: list[tuple[str, str]]) -> None:
    """A document with real length: many paragraphs per section."""
    from docx import Document

    doc = Document()
    for kind, text in body:
        if kind == "h":
            doc.add_heading(text, level=1)
        else:
            doc.add_paragraph(text)
    doc.save(path)


def _placeholder() -> str:
    return "Filler paragraph so the section has enough content to push a figure off a page. " * 3


def _body_with_image_marker() -> list[tuple[str, str]]:
    body: list[tuple[str, str]] = [("h", "1. Introduction")]
    body += [("p", _placeholder()) for _ in range(12)]
    body.append(("h", "2. Background"))
    body += [("p", _placeholder()) for _ in range(12)]
    body.append(("h", "3. Method"))
    body += [("p", _placeholder()) for _ in range(4)]
    body.append(("p", "@@FIGURE@@"))
    body += [("p", _placeholder()) for _ in range(4)]
    body.append(("h", "4. Results"))
    body += [("p", _placeholder()) for _ in range(12)]
    body.append(("h", "5. Conclusion"))
    body += [("p", _placeholder()) for _ in range(12)]
    return body


def _section_of_image(path: Path) -> tuple[str, bool]:
    """Walk the DOCX body in order and report the heading the image sits under.

    Reads the saved file rather than trusting the builder, and returns whether
    the picture is inline (the only thing python-docx can produce).
    """
    import docx
    from docx.oxml.ns import qn

    doc = docx.Document(str(path))
    current = ""
    inline = False
    for para in doc.paragraphs:
        if para.style.name.startswith("Heading"):
            current = para.text
            continue
        blips = para._p.findall(".//" + qn("a:blip"))
        if not blips:
            continue
        inline = para._p.findall(".//" + qn("wp:inline")) != []
        return current, inline
    raise AssertionError("no picture found in the document")


def test_an_image_lands_inside_the_section_it_was_asked_for(tmp_path: Path, png: Path) -> None:
    """The long-document case: 60+ paragraphs, and the figure must be in 3, not 2 or 4."""
    from docx import Document
    from docx.shared import Inches

    body = _body_with_image_marker()
    out = tmp_path / "long.docx"
    _long_docx(out, [item for item in body if item[1] != "@@FIGURE@@"])

    # Rebuild, inserting the picture at the marked paragraph.
    doc = Document()
    for kind, text in body:
        if text == "@@FIGURE@@":
            doc.add_paragraph().add_run().add_picture(str(png), width=Inches(4))
        elif kind == "h":
            doc.add_heading(text, level=1)
        else:
            doc.add_paragraph(text)
    doc.save(out)

    section, inline = _section_of_image(out)
    assert section == "3. Method", f"figure landed in {section!r}"
    assert inline, "python-docx can only place inline images; an anchor means raw XML"


def test_the_image_is_a_real_picture_not_a_blank_frame(tmp_path: Path, png: Path) -> None:
    import zipfile

    from docx import Document
    from docx.shared import Inches

    out = tmp_path / "figure.docx"
    doc = Document()
    doc.add_paragraph("Before")
    doc.add_picture(str(png), width=Inches(3))
    doc.save(out)

    with zipfile.ZipFile(out) as archive:
        media = [n for n in archive.namelist() if "/media/" in n]
        assert media
        embedded = Image.open(io.BytesIO(archive.read(media[0])))
    embedded.load()
    assert embedded.size == (640, 360)
    # A rendered frame, not one flat colour.
    assert len({embedded.getpixel(p) for p in ((10, 10), (400, 10), (400, 300))}) > 1


def test_width_only_is_what_keeps_the_aspect_ratio(tmp_path: Path, png: Path) -> None:
    """The squashed-image failure, caught by measurement rather than eyeballing."""
    import zipfile

    from docx import Document
    from docx.shared import Inches

    good = tmp_path / "good.docx"
    doc = Document()
    doc.add_picture(str(png), width=Inches(4))
    doc.save(good)

    from docx.oxml.ns import qn
    import docx

    para = docx.Document(str(good)).paragraphs[0]
    extent = para._p.find(".//" + qn("wp:extent"))
    cx, cy = int(extent.get("cx")), int(extent.get("cy"))
    assert abs((cx / cy) - (640 / 360)) < 0.01, "aspect ratio changed"


def test_a_table_cell_pins_a_picture_beside_text(tmp_path: Path, png: Path) -> None:
    """In DOCX, a table is how you get a picture where a flow would not put it."""
    import docx
    from docx import Document
    from docx.shared import Inches

    out = tmp_path / "column.docx"
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Caption text beside the figure, which is what a two-column layout is for."
    table.cell(0, 1).paragraphs[0].add_run().add_picture(str(png), width=Inches(2.5))
    doc.save(out)

    reopened = docx.Document(str(out))
    cell = reopened.tables[0].cell(0, 1)
    assert cell.paragraphs[0].runs[0]._r.findall(
        ".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
    )


def test_pptx_places_an_image_at_exact_coordinates(tmp_path: Path, png: Path) -> None:
    """A slide is a fixed canvas, so the inches survive to the file."""
    from pptx import Presentation
    from pptx.util import Inches

    out = tmp_path / "slide.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    left, top, width = Inches(2), Inches(1.5), Inches(4)
    slide.shapes.add_picture(str(png), left, top, width=width)
    prs.save(out)

    reopened = Presentation(str(out))
    picture = next(s for s in reopened.slides[0].shapes if s.shape_type == 13)
    assert (picture.left, picture.top, picture.width) == (left, top, width)
    # Height follows the aspect ratio: 4in wide at 16:9 is 2.25in.
    assert abs(picture.height - Inches(2.25)) < Inches(0.02)
