#!/usr/bin/env python3
"""Document work inside the sandbox: convert, create, image, chart, insert, job.

Bootstrapped into the sandbox by ``nanobot/agent/tools/document_forge.py`` and
driven from there, the same way ``engineering_draw_cli.py`` is. Every action
prints one JSON object as its last stdout line. Heavy actions take
``--background``: the CLI forks itself detached, writes a job record under
``~/.document_forge/jobs``, and the record turns into ``done`` or ``failed`` when
the work ends. ``job --job-id ID`` reads it.

Python libraries are installed with ``install`` into ``~/.document_forge/site``
(``pip --target``), so the sandbox image does not need them preinstalled. Office
conversion needs LibreOffice and PDF rasterising needs poppler; ``install``
attempts those with apt when the sandbox user is root or has passwordless sudo,
and reports plainly when it cannot.
"""

from __future__ import annotations

import argparse
import csv
import importlib
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import traceback
import urllib.request
import uuid
from pathlib import Path
from typing import Any

CLI_VERSION = "document-forge-2026.1"

HOME = Path(os.environ.get("DOCUMENT_FORGE_HOME") or (Path.home() / ".document_forge"))
SITE = HOME / "site"
JOBS = HOME / "jobs"
DEFAULT_OUT = os.environ.get("DOCUMENT_FORGE_DIR", "~/document_outputs")
MAX_DOWNLOAD_BYTES = 25 * 1024 * 1024
MAX_PDF_PAGES_RASTERISED = 50
os.environ.setdefault("MPLCONFIGDIR", str(HOME / "mplconfig"))

PIP_PACKAGES = ["python-docx", "python-pptx", "openpyxl", "matplotlib", "pillow", "markdown"]
#: (import name, distribution name) for doctor.
MODULES = [
    ("docx", "python-docx"),
    ("pptx", "python-pptx"),
    ("openpyxl", "openpyxl"),
    ("matplotlib", "matplotlib"),
    ("PIL", "pillow"),
    ("markdown", "markdown"),
]
APT_PACKAGES = [
    "libreoffice-writer-nogui",
    "libreoffice-calc-nogui",
    "libreoffice-impress-nogui",
    "poppler-utils",
    "fonts-liberation",
    "fonts-dejavu-core",
]
IMAGE_FORMATS = {
    "png": "PNG",
    "jpg": "JPEG",
    "jpeg": "JPEG",
    "webp": "WEBP",
    "gif": "GIF",
    "bmp": "BMP",
    "tif": "TIFF",
    "tiff": "TIFF",
}
DOCX_EMBEDDABLE = {"png", "jpg", "jpeg", "gif", "bmp", "tif", "tiff"}
OFFICE_INPUTS = {
    "doc",
    "docx",
    "odt",
    "rtf",
    "txt",
    "html",
    "htm",
    "ppt",
    "pptx",
    "odp",
    "xls",
    "xlsx",
    "ods",
    "csv",
}
PALETTE = ["#1F4E79", "#C55A11", "#548235", "#7030A0", "#BF9000", "#2E75B6", "#7F7F7F"]
_INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*)")
_URL = re.compile(r"^https?://", re.IGNORECASE)


class CliError(Exception):
    def __init__(self, code: str, message: str, next_step: str | None = None) -> None:
        super().__init__(message)
        self.code = code
        self.message = message
        self.next_step = next_step


# --------------------------------------------------------------------------- #
# Plumbing
# --------------------------------------------------------------------------- #
def _emit(payload: dict[str, Any]) -> None:
    print(json.dumps(payload, default=str, ensure_ascii=False))
    sys.stdout.flush()


def _now() -> str:
    return time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())


def _ensure_site_on_path() -> None:
    if SITE.is_dir() and str(SITE) not in sys.path:
        sys.path.insert(0, str(SITE))
    importlib.invalidate_caches()


def _module(name: str, pip_name: str) -> Any:
    _ensure_site_on_path()
    try:
        return importlib.import_module(name)
    except ImportError as exc:
        raise CliError(
            "missing_dependency",
            f"Python package '{pip_name}' is not installed ({exc}).",
            "Run document_forge action='install' with background=true, poll action='job' "
            "until it reports done, then retry.",
        ) from exc


def _safe_name(value: Any, default: str) -> str:
    text = re.sub(r"[^A-Za-z0-9._-]+", "-", str(value or "")).strip("-.")
    return (text or default)[:120]


def _out_dir(value: str | None) -> Path:
    path = Path(os.path.expanduser(value or DEFAULT_OUT))
    path.mkdir(parents=True, exist_ok=True)
    return path


def _unique(path: Path) -> Path:
    if not path.exists():
        return path
    for index in range(2, 10000):
        candidate = path.with_name(f"{path.stem}-{index}{path.suffix}")
        if not candidate.exists():
            return candidate
    raise CliError("too_many_files", f"Too many files named like {path.name}.")


def _input_file(value: str | None, label: str = "file") -> Path:
    if not value:
        raise CliError("missing_argument", f"This action needs {label}.")
    path = Path(os.path.expanduser(value))
    if not path.is_file():
        raise CliError("file_not_found", f"No file at {path}.")
    return path


def _load_spec(value: Any) -> Any:
    if value is None:
        raise CliError(
            "missing_spec", "This action needs a spec: a JSON object or a path to a .json file."
        )
    if isinstance(value, (dict, list)):
        return value
    text = str(value)
    path = Path(os.path.expanduser(text))
    if not text.lstrip().startswith(("{", "[")) and path.is_file():
        text = path.read_text(encoding="utf-8")
    try:
        return json.loads(text)
    except json.JSONDecodeError as exc:
        raise CliError("bad_spec", f"The spec is not valid JSON: {exc}.") from exc


def _fetch_bytes(url: str) -> bytes:
    if not _URL.match(url):
        raise CliError("bad_url", f"Only http(s) URLs can be fetched, got {url!r}.")
    request = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 (document-forge)"})
    try:
        with urllib.request.urlopen(request, timeout=30) as response:
            data = response.read(MAX_DOWNLOAD_BYTES + 1)
    except Exception as exc:  # noqa: BLE001 - reported to the caller
        raise CliError("fetch_failed", f"Could not download {url}: {exc}.") from exc
    if len(data) > MAX_DOWNLOAD_BYTES:
        raise CliError(
            "too_large", f"{url} is larger than {MAX_DOWNLOAD_BYTES // (1024 * 1024)} MB."
        )
    return data


def _pil() -> Any:
    return _module("PIL.Image", "pillow")


def _open_image(src: str) -> tuple[Any, str]:
    """Open a raster image from a path or URL. Returns (image, original format)."""
    pil_image = _pil()
    image_ops = importlib.import_module("PIL.ImageOps")
    if _URL.match(str(src)):
        source: Any = io.BytesIO(_fetch_bytes(str(src)))
    else:
        path = Path(os.path.expanduser(str(src)))
        if not path.is_file():
            raise CliError("file_not_found", f"No image at {path}.")
        source = path
    try:
        image = pil_image.open(source)
        fmt = (image.format or "PNG").upper()
        image.load()
    except Exception as exc:  # noqa: BLE001
        raise CliError(
            "not_an_image",
            f"{src} is not a raster image PIL can read ({exc}). SVG and PDF are not raster "
            "images: convert them to PNG first.",
        ) from exc
    return image_ops.exif_transpose(image), fmt


def _parse_crop(value: Any) -> Any:
    """Normalise a crop: None, an 'W:H' aspect string, or an (x, y, w, h) tuple."""
    if value is None or value == "":
        return None
    if isinstance(value, dict):
        try:
            return tuple(int(value[k]) for k in ("x", "y", "w", "h"))
        except (KeyError, TypeError, ValueError) as exc:
            raise CliError("bad_crop", "crop as an object needs integer x, y, w and h.") from exc
    if isinstance(value, (list, tuple)):
        if len(value) != 4:
            raise CliError("bad_crop", "crop as a list needs exactly four numbers: x, y, w, h.")
        return tuple(int(v) for v in value)
    text = str(value).strip()
    if ":" in text:
        try:
            aw, ah = (float(part) for part in text.split(":"))
        except ValueError as exc:
            raise CliError("bad_crop", f"crop aspect {text!r} must look like 16:9.") from exc
        if aw <= 0 or ah <= 0:
            raise CliError("bad_crop", f"crop aspect {text!r} must be positive.")
        return text
    try:
        parts = [int(float(part)) for part in text.split(",")]
    except ValueError as exc:
        raise CliError(
            "bad_crop", f"crop {text!r} must be 'x,y,w,h' pixels or an aspect like 16:9."
        ) from exc
    if len(parts) != 4:
        raise CliError(
            "bad_crop", f"crop {text!r} must be 'x,y,w,h' pixels or an aspect like 16:9."
        )
    return tuple(parts)


def _crop(image: Any, crop: Any) -> Any:
    if crop is None:
        return image
    width, height = image.size
    if isinstance(crop, str):
        aw, ah = (float(part) for part in crop.split(":"))
        target = aw / ah
        if width / height > target:
            w, h = round(height * target), height
        else:
            w, h = width, round(width / target)
        x, y = (width - w) // 2, (height - h) // 2
    else:
        x, y, w, h = crop
    if w <= 0 or h <= 0 or x < 0 or y < 0 or x + w > width or y + h > height:
        raise CliError(
            "bad_crop",
            f"crop x={x} y={y} w={w} h={h} lies outside the {width}x{height} image.",
            "Read the image size first (action='image' with op='convert' reports it) and crop inside it.",
        )
    return image.crop((x, y, x + w, y + h))


def _save_image(image: Any, path: Path, fmt: str | None = None) -> Path:
    pil_image = _pil()
    ext = (fmt or path.suffix.lstrip(".")).lower()
    pil_format = IMAGE_FORMATS.get(ext)
    if pil_format is None:
        raise CliError(
            "bad_format", f"Unsupported image format {ext!r}. Use one of {sorted(IMAGE_FORMATS)}."
        )
    if pil_format == "JPEG":
        if image.mode in ("RGBA", "LA", "P"):
            rgba = image.convert("RGBA")
            flat = pil_image.new("RGB", rgba.size, "white")
            flat.paste(rgba, mask=rgba.getchannel("A"))
            image = flat
        elif image.mode != "RGB":
            image = image.convert("RGB")
        image.save(path, format=pil_format, quality=92)
    else:
        image.save(path, format=pil_format)
    return path


def _materialise(src: str, crop: Any, workdir: Path) -> Path:
    """Return a file path python-docx can embed, fetching or cropping if asked."""
    crop = _parse_crop(crop)
    if crop is None and not _URL.match(str(src)):
        path = Path(os.path.expanduser(str(src)))
        if path.is_file() and path.suffix.lower().lstrip(".") in DOCX_EMBEDDABLE:
            return path
    image, _fmt = _open_image(src)
    image = _crop(image, crop)
    return _save_image(image, workdir / f"img-{uuid.uuid4().hex[:8]}.png", "png")


def _pdf_pages(pdf: Path) -> int | None:
    exe = shutil.which("pdfinfo")
    if not exe:
        return None
    proc = subprocess.run([exe, str(pdf)], capture_output=True, text=True, timeout=60)
    match = re.search(r"^Pages:\s+(\d+)", proc.stdout, re.MULTILINE)
    return int(match.group(1)) if match else None


# --------------------------------------------------------------------------- #
# LibreOffice and poppler
# --------------------------------------------------------------------------- #
def _office(src: Path, to: str, outdir: Path, timeout: int = 600) -> Path:
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        raise CliError(
            "no_office",
            "LibreOffice is not installed in this sandbox.",
            "Run document_forge action='install' with background=true (it installs LibreOffice "
            "when apt is available), poll action='job', then retry.",
        )
    profile = Path(tempfile.mkdtemp(prefix="lo-profile-"))
    cmd = [
        exe,
        f"-env:UserInstallation=file://{profile}",
        "--headless",
        "--norestore",
        "--nolockcheck",
        "--convert-to",
        to,
        "--outdir",
        str(outdir),
        str(src),
    ]
    try:
        proc = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
    except subprocess.TimeoutExpired as exc:
        raise CliError(
            "convert_timeout", f"LibreOffice did not finish {src.name} within {timeout}s."
        ) from exc
    finally:
        shutil.rmtree(profile, ignore_errors=True)
    produced = outdir / f"{src.stem}.{to.split(':')[0]}"
    if not produced.exists():
        detail = (proc.stderr or proc.stdout or "").strip()[-400:]
        raise CliError(
            "convert_failed", f"LibreOffice produced no {to} file for {src.name}. {detail}"
        )
    return produced


def _pdf_to_images(pdf: Path, out_dir: Path, stem: str, fmt: str, dpi: int = 150) -> list[str]:
    exe = shutil.which("pdftoppm")
    if not exe:
        raise CliError(
            "missing_binary",
            "pdftoppm is not installed (poppler-utils).",
            "Run document_forge action='install' with background=true, then retry.",
        )
    ext = "jpg" if fmt in ("jpg", "jpeg") else "png"
    flag = "-jpeg" if ext == "jpg" else "-png"
    with tempfile.TemporaryDirectory(prefix="pdfpages-") as tmp:
        proc = subprocess.run(
            [exe, flag, "-r", str(dpi), str(pdf), str(Path(tmp) / "page")],
            capture_output=True,
            text=True,
            timeout=600,
        )
        if proc.returncode != 0:
            raise CliError("render_failed", f"pdftoppm failed: {proc.stderr.strip()[-300:]}")
        pages = sorted(
            Path(tmp).glob(f"page-*.{ext}"),
            key=lambda p: int(re.search(r"-(\d+)\.", p.name).group(1)),
        )
        if not pages:
            raise CliError("render_failed", "pdftoppm produced no pages.")
        if len(pages) > MAX_PDF_PAGES_RASTERISED:
            raise CliError(
                "too_many_pages",
                f"{len(pages)} pages is more than the {MAX_PDF_PAGES_RASTERISED} rasterised at once.",
                "Convert a smaller page range or split the document.",
            )
        written: list[str] = []
        for index, page in enumerate(pages, start=1):
            target_name = f"{stem}.{ext}" if len(pages) == 1 else f"{stem}-page-{index}.{ext}"
            target = _unique(out_dir / target_name)
            shutil.move(str(page), target)
            written.append(str(target))
    return written


# --------------------------------------------------------------------------- #
# DOCX building
# --------------------------------------------------------------------------- #
def _rgb(hex_color: str) -> Any:
    docx_shared = _module("docx.shared", "python-docx")
    return docx_shared.RGBColor.from_string(hex_color.lstrip("#").upper())


def _force_font(style: Any, font: str) -> None:
    """Set a style's font and drop theme-font overrides that would win over it."""
    from docx.oxml.ns import qn

    style.font.name = font
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        return
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)
    rfonts.set(qn("w:eastAsia"), font)


def _shade(cell: Any, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _bottom_rule(paragraph: Any, color: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def _page_number(paragraph: Any) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def add(child: Any) -> None:
        run = OxmlElement("w:r")
        run.append(child)
        paragraph._p.append(run)

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    add(begin)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    add(instr)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    add(separate)
    text = OxmlElement("w:t")
    text.text = "1"
    add(text)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    add(end)


def _rich(
    paragraph: Any,
    text: Any,
    *,
    size: float | None = None,
    color: str | None = None,
    bold: bool = False,
    italic: bool = False,
    font: str | None = None,
) -> None:
    """Add text with **bold**, *italic* and `code` spans."""
    for part in _INLINE.split(str(text)):
        if not part:
            continue
        run_bold, run_italic, mono = bold, italic, False
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            part, run_bold = part[2:-2], True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            part, mono = part[1:-1], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, run_italic = part[1:-1], True
        run = paragraph.add_run(part)
        run.bold = run_bold
        run.italic = run_italic
        if mono:
            run.font.name = "Consolas"
        elif font:
            run.font.name = font
        if size:
            run.font.size = _pt(size)
        if color:
            run.font.color.rgb = _rgb(color)


def _pt(value: float) -> Any:
    return _module("docx.shared", "python-docx").Pt(value)


def _cm(value: float) -> Any:
    return _module("docx.shared", "python-docx").Cm(value)


def _align(name: str | None) -> Any:
    enum = importlib.import_module("docx.enum.text")
    table = {
        "left": enum.WD_ALIGN_PARAGRAPH.LEFT,
        "center": enum.WD_ALIGN_PARAGRAPH.CENTER,
        "right": enum.WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": enum.WD_ALIGN_PARAGRAPH.JUSTIFY,
    }
    return table.get(str(name or "left").lower(), enum.WD_ALIGN_PARAGRAPH.LEFT)


def _usable_cm(document: Any) -> float:
    section = document.sections[0]
    return (section.page_width - section.left_margin - section.right_margin) / 360000.0


def _picture(
    document: Any,
    path: Path,
    width_cm: float,
    align: str | None,
    caption: str | None,
    font: str | None,
) -> None:
    width_cm = min(float(width_cm), _usable_cm(document))
    paragraph = document.add_paragraph()
    paragraph.alignment = _align(align or "center")
    paragraph.add_run().add_picture(str(path), width=_cm(width_cm))
    if caption:
        note = document.add_paragraph()
        note.alignment = _align("center")
        _rich(note, caption, size=9, italic=True, font=font)


def _table(document: Any, block: dict[str, Any], primary: str, font: str | None) -> None:
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    headers = [str(h) for h in block.get("headers") or []]
    rows = [[str(cell) for cell in row] for row in block.get("rows") or []]
    columns = len(headers) or (max((len(r) for r in rows), default=0))
    if columns == 0:
        raise CliError("bad_table", "A table needs headers or rows.")
    if block.get("caption"):
        caption = document.add_paragraph()
        _rich(caption, block["caption"], size=9.5, bold=True, font=font)
    table = document.add_table(rows=0, cols=columns)
    table.style = document.styles["Table Grid"]
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    if headers:
        row = table.add_row()
        tr_pr = row._tr.get_or_add_trPr()
        header_flag = OxmlElement("w:tblHeader")
        header_flag.set(qn("w:val"), "true")
        tr_pr.append(header_flag)
        for index in range(columns):
            cell = row.cells[index]
            _shade(cell, primary)
            paragraph = cell.paragraphs[0]
            text = headers[index] if index < len(headers) else ""
            _rich(paragraph, text, size=9.5, bold=True, color="FFFFFF", font=font)
    for number, values in enumerate(rows):
        row = table.add_row()
        for index in range(columns):
            cell = row.cells[index]
            if number % 2 == 1:
                _shade(cell, "F2F2F2")
            text = values[index] if index < len(values) else ""
            _rich(cell.paragraphs[0], text, size=9.5, font=font)
    widths = block.get("col_widths_cm")
    if widths:
        for row in table.rows:
            for index, cell in enumerate(row.cells):
                if index < len(widths):
                    cell.width = _cm(float(widths[index]))
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = _pt(2)


def _render_chart(spec: dict[str, Any], out: Path) -> Path:
    mpl = _module("matplotlib", "matplotlib")
    mpl.use("Agg")
    plt = importlib.import_module("matplotlib.pyplot")
    kind = str(spec.get("kind", "bar")).lower()
    labels = [str(x) for x in spec.get("labels") or []]
    series = spec.get("series") or {}
    if kind not in ("bar", "barh", "line", "pie"):
        raise CliError("bad_chart", f"Unknown chart kind {kind!r}. Use bar, barh, line or pie.")
    if not labels or not isinstance(series, dict) or not series:
        raise CliError("bad_chart", "A chart needs labels and at least one named series.")
    values_by_name: dict[str, list[float]] = {}
    for name, values in series.items():
        if len(values) != len(labels):
            raise CliError(
                "bad_chart", f"Series '{name}' has {len(values)} values for {len(labels)} labels."
            )
        try:
            values_by_name[str(name)] = [float(v) for v in values]
        except (TypeError, ValueError) as exc:
            raise CliError("bad_chart", f"Series '{name}' has a non-numeric value.") from exc

    width, height = spec.get("size_in") or [8, 4.8]
    fig, ax = plt.subplots(figsize=(float(width), float(height)), dpi=200)
    try:
        if kind == "pie":
            first_name, first_values = next(iter(values_by_name.items()))
            ax.pie(
                first_values,
                labels=labels,
                colors=PALETTE[: len(labels)] or None,
                autopct="%1.1f%%",
                startangle=90,
                counterclock=False,
            )
            ax.axis("equal")
        else:
            count = len(values_by_name)
            positions = list(range(len(labels)))
            bar_width = 0.8 / count
            for index, (name, values) in enumerate(values_by_name.items()):
                color = PALETTE[index % len(PALETTE)]
                offset = [p - 0.4 + bar_width * (index + 0.5) for p in positions]
                if kind == "line":
                    ax.plot(labels, values, marker="o", linewidth=2, color=color, label=name)
                elif kind == "barh":
                    ax.barh(offset, values, height=bar_width, color=color, label=name)
                else:
                    ax.bar(offset, values, width=bar_width, color=color, label=name)
            if kind == "barh":
                ax.set_yticks(positions)
                ax.set_yticklabels(labels)
                ax.grid(axis="x", alpha=0.3)
            else:
                if kind == "bar":
                    ax.set_xticks(positions)
                    ax.set_xticklabels(
                        labels,
                        rotation=30 if len(labels) > 8 else 0,
                        ha="right" if len(labels) > 8 else "center",
                    )
                ax.grid(axis="y", alpha=0.3)
            ax.spines["top"].set_visible(False)
            ax.spines["right"].set_visible(False)
            ax.set_xlabel(str(spec.get("xlabel", "")))
            ax.set_ylabel(str(spec.get("ylabel", "")))
            if count > 1:
                ax.legend(frameon=False)
        ax.set_title(str(spec.get("title", "")), fontsize=13, weight="bold")
        fig.tight_layout()
        fig.savefig(out, dpi=200)
    finally:
        plt.close(fig)
    return out


# --- Markdown -> blocks ----------------------------------------------------- #
def _split_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _md_to_blocks(text: str) -> list[dict[str, Any]]:
    blocks: list[dict[str, Any]] = []
    lines = text.splitlines()
    pending: list[str] = []

    def flush() -> None:
        if pending:
            blocks.append({"type": "paragraph", "text": " ".join(p.strip() for p in pending)})
            pending.clear()

    index = 0
    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()
        if stripped.startswith("```"):
            flush()
            code: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code.append(lines[index])
                index += 1
            blocks.append({"type": "code", "text": "\n".join(code)})
            index += 1
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            flush()
            blocks.append(
                {
                    "type": "heading",
                    "level": min(len(heading.group(1)), 3),
                    "text": heading.group(2).strip(),
                }
            )
            index += 1
            continue
        if (
            stripped.startswith("|")
            and index + 1 < len(lines)
            and re.match(r"^\|?\s*:?-{2,}", lines[index + 1].strip())
        ):
            flush()
            headers = _split_row(stripped)
            index += 2
            rows: list[list[str]] = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(_split_row(lines[index].strip()))
                index += 1
            blocks.append({"type": "table", "headers": headers, "rows": rows})
            continue
        image = re.match(r"^!\[([^\]]*)\]\(([^)\s]+)\)\s*$", stripped)
        if image:
            flush()
            source = image.group(2)
            block: dict[str, Any] = {"type": "image", "caption": image.group(1) or None}
            block["url" if _URL.match(source) else "path"] = source
            blocks.append(block)
            index += 1
            continue
        if re.match(r"^[-*+]\s+", stripped):
            flush()
            items: list[str] = []
            while index < len(lines) and re.match(r"^\s*[-*+]\s+", lines[index]):
                items.append(re.sub(r"^\s*[-*+]\s+", "", lines[index]).strip())
                index += 1
            blocks.append({"type": "bullets", "items": items})
            continue
        if re.match(r"^\d+[.)]\s+", stripped):
            flush()
            items = []
            while index < len(lines) and re.match(r"^\s*\d+[.)]\s+", lines[index]):
                items.append(re.sub(r"^\s*\d+[.)]\s+", "", lines[index]).strip())
                index += 1
            blocks.append({"type": "numbered", "items": items})
            continue
        if stripped == "":
            flush()
        elif re.match(r"^(-{3,}|\*{3,}|_{3,})$", stripped):
            flush()
        elif stripped.startswith(">"):
            flush()
            blocks.append(
                {"type": "paragraph", "text": stripped.lstrip("> ").strip(), "italic": True}
            )
        else:
            pending.append(stripped)
        index += 1
    flush()
    return blocks


def _build_docx(spec: dict[str, Any], out: Path, workdir: Path) -> dict[str, Any]:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    if not isinstance(spec, dict):
        raise CliError("bad_spec", "A document spec is a JSON object with a 'blocks' list.")
    document = Document()
    font = str(spec.get("font") or "Calibri")
    body_pt = float(spec.get("body_size_pt") or 11)
    primary = str(spec.get("primary_color") or "1F4E79").lstrip("#")

    normal = document.styles["Normal"]
    _force_font(normal, font)
    normal.font.size = _pt(body_pt)
    normal.paragraph_format.space_after = _pt(6)
    for level, size in ((1, 16), (2, 13), (3, 11.5)):
        style = document.styles[f"Heading {level}"]
        _force_font(style, font)
        style.font.size = _pt(size)
        style.font.bold = True
        style.font.color.rgb = _rgb(primary)
        style.paragraph_format.space_before = _pt(14 if level == 1 else 10)
        style.paragraph_format.space_after = _pt(6)

    section = document.sections[0]
    if str(spec.get("page_size") or "A4").upper() == "LETTER":
        section.page_width, section.page_height = _cm(21.59), _cm(27.94)
    else:
        section.page_width, section.page_height = _cm(21.0), _cm(29.7)
    if str(spec.get("orientation") or "").lower() == "landscape":
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    margin = float(spec.get("margin_cm") or 2.2)
    section.left_margin = section.right_margin = _cm(margin)
    section.top_margin = section.bottom_margin = _cm(margin)

    if spec.get("header_text"):
        head = section.header.paragraphs[0]
        head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _rich(head, spec["header_text"], size=8.5, color="7F7F7F", font=font)
    footer_text = spec.get("footer_text")
    if footer_text or spec.get("page_numbers", True):
        foot = section.footer.paragraphs[0]
        foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if footer_text:
            _rich(foot, f"{footer_text}   |   Page ", size=8.5, color="7F7F7F", font=font)
        else:
            _rich(foot, "Page ", size=8.5, color="7F7F7F", font=font)
        if spec.get("page_numbers", True):
            _page_number(foot)

    logo = spec.get("logo")
    if logo:
        if isinstance(logo, str):
            logo = {"path": logo}
        logo_path = _materialise(logo.get("url") or logo.get("path"), logo.get("crop"), workdir)
        _picture(
            document,
            logo_path,
            float(logo.get("width_cm") or 3.5),
            logo.get("align") or "left",
            None,
            font,
        )

    if spec.get("title"):
        title = document.add_paragraph()
        _rich(title, spec["title"], size=24, bold=True, color=primary, font=font)
        title.paragraph_format.space_after = _pt(4)
    if spec.get("subtitle"):
        sub = document.add_paragraph()
        _rich(sub, spec["subtitle"], size=13, color="595959", font=font)
    meta_parts = [str(spec[k]) for k in ("author", "date") if spec.get(k)]
    if meta_parts:
        meta = document.add_paragraph()
        _rich(meta, "  ·  ".join(meta_parts), size=10, color="7F7F7F", font=font)
        if spec.get("title") or spec.get("subtitle"):
            _bottom_rule(meta, primary)

    blocks = spec.get("blocks") or []
    if not isinstance(blocks, list):
        raise CliError("bad_spec", "'blocks' must be a list.")
    for position, block in enumerate(blocks, start=1):
        if not isinstance(block, dict):
            raise CliError("bad_block", f"Block {position} is not an object.")
        kind = str(block.get("type", "paragraph")).lower()
        if kind == "heading":
            level = min(max(int(block.get("level", 1)), 1), 3)
            paragraph = document.add_paragraph(style=f"Heading {level}")
            _rich(paragraph, block.get("text", ""), font=font)
        elif kind == "paragraph":
            paragraph = document.add_paragraph()
            paragraph.alignment = _align(block.get("align"))
            _rich(
                paragraph,
                block.get("text", ""),
                size=block.get("size_pt"),
                font=font,
                italic=bool(block.get("italic")),
            )
        elif kind in ("bullets", "numbered"):
            items = block.get("items") or []
            for number, item in enumerate(items, start=1):
                paragraph = document.add_paragraph(
                    style="List Bullet" if kind == "bullets" else None
                )
                if kind == "numbered":
                    paragraph.paragraph_format.left_indent = _cm(0.63)
                    paragraph.paragraph_format.first_line_indent = _cm(-0.63)
                    paragraph.add_run(f"{number}.\t")
                _rich(paragraph, item, font=font)
        elif kind == "table":
            _table(document, block, primary, font)
        elif kind == "image":
            source = block.get("url") or block.get("path")
            if not source:
                raise CliError("bad_block", f"Image block {position} needs 'path' or 'url'.")
            path = _materialise(source, block.get("crop"), workdir)
            _picture(
                document,
                path,
                float(block.get("width_cm") or 14),
                block.get("align"),
                block.get("caption"),
                font,
            )
        elif kind == "chart":
            chart_png = _render_chart(block, workdir / f"chart-{uuid.uuid4().hex[:8]}.png")
            _picture(
                document,
                chart_png,
                float(block.get("width_cm") or 15),
                "center",
                block.get("caption"),
                font,
            )
        elif kind == "code":
            paragraph = document.add_paragraph()
            _rich(paragraph, block.get("text", ""), size=9, font="Consolas")
        elif kind == "page_break":
            document.add_page_break()
        else:
            raise CliError("bad_block", f"Unknown block type {kind!r} at position {position}.")

    document.core_properties.title = str(spec.get("title") or out.stem)
    if spec.get("author"):
        document.core_properties.author = str(spec["author"])
    document.save(out)
    return _docx_stats(out)


def _docx_stats(path: Path) -> dict[str, Any]:
    from docx import Document

    reopened = Document(path)
    return {
        "paragraphs": len(reopened.paragraphs),
        "tables": len(reopened.tables),
        "images": len(reopened.inline_shapes),
    }


def _docx_to_md(path: Path) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(path)
    parts: list[str] = []
    for child in document.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            style = paragraph.style.name if paragraph.style is not None else ""
            heading = re.match(r"Heading (\d)", style)
            if heading:
                parts.append("#" * min(int(heading.group(1)), 6) + " " + text)
            elif style.startswith("List Bullet"):
                parts.append("- " + text)
            elif style.startswith("List Number"):
                parts.append("1. " + text)
            else:
                parts.append(text)
        elif tag == "tbl":
            table = Table(child, document)
            rows = [
                [cell.text.strip().replace("\n", " ") for cell in row.cells] for row in table.rows
            ]
            if not rows:
                continue
            width = max(len(r) for r in rows)
            rows = [r + [""] * (width - len(r)) for r in rows]
            lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
            lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
            parts.append("\n".join(lines))
    return "\n\n".join(parts) + "\n"


# --------------------------------------------------------------------------- #
# Actions
# --------------------------------------------------------------------------- #
def do_convert(args: argparse.Namespace) -> dict[str, Any]:
    src = _input_file(args.file)
    target = str(args.to or "").lower().lstrip(".")
    if not target:
        raise CliError(
            "missing_argument", "convert needs 'to', the target format, e.g. pdf or docx."
        )
    source = src.suffix.lower().lstrip(".")
    if source == target:
        raise CliError("same_format", f"The file is already {target}.")
    out_dir = _out_dir(args.out_dir)
    stem = _safe_name(args.name, src.stem)
    workdir = Path(tempfile.mkdtemp(prefix="docconv-"))
    try:
        if source in ("md", "markdown"):
            blocks = _md_to_blocks(src.read_text(encoding="utf-8"))
            if target == "docx":
                outputs = [_write_docx_output({"blocks": blocks}, out_dir, stem, workdir)]
            elif target == "pdf":
                docx_tmp = workdir / f"{stem}.docx"
                _build_docx({"blocks": blocks}, docx_tmp, workdir)
                outputs = [_finish_pdf(docx_tmp, out_dir, stem, workdir)]
            elif target == "html":
                markdown = _module("markdown", "markdown")
                body = markdown.markdown(
                    src.read_text(encoding="utf-8"), extensions=["tables", "fenced_code"]
                )
                target_path = _unique(out_dir / f"{stem}.html")
                target_path.write_text(
                    "<!doctype html><meta charset='utf-8'><style>body{font-family:Calibri,Arial,sans-serif;"
                    "max-width:820px;margin:2rem auto;line-height:1.5}table{border-collapse:collapse}"
                    "td,th{border:1px solid #bbb;padding:4px 8px}th{background:#1F4E79;color:#fff}</style>"
                    f"{body}",
                    encoding="utf-8",
                )
                outputs = [_output(target_path, "html")]
            else:
                raise CliError(
                    "unsupported_conversion", f"md can convert to docx, pdf or html, not {target}."
                )
        elif source in DOCX_EMBEDDABLE or source == "webp":
            if target in IMAGE_FORMATS:
                image, _fmt = _open_image(str(src))
                target_path = _unique(out_dir / f"{stem}.{target}")
                _save_image(image, target_path, target)
                outputs = [_output(target_path, target, image=image)]
            elif target == "pdf":
                image, _fmt = _open_image(str(src))
                target_path = _unique(out_dir / f"{stem}.pdf")
                if image.mode != "RGB":
                    image = image.convert("RGB")
                image.save(target_path, "PDF", resolution=150.0)
                outputs = [_output(target_path, "pdf")]
            else:
                raise CliError(
                    "unsupported_conversion",
                    f"Images convert to pdf or image formats, not {target}.",
                )
        elif source == "pdf":
            if target == "txt":
                exe = shutil.which("pdftotext")
                if not exe:
                    raise CliError(
                        "missing_binary",
                        "pdftotext is not installed (poppler-utils).",
                        "Run document_forge action='install' with background=true, then retry.",
                    )
                target_path = _unique(out_dir / f"{stem}.txt")
                subprocess.run(
                    [exe, "-layout", str(src), str(target_path)], check=True, timeout=600
                )
                outputs = [_output(target_path, "txt")]
            elif target in ("png", "jpg", "jpeg"):
                paths = _pdf_to_images(src, out_dir, stem, target)
                outputs = [_output(Path(p), target) for p in paths]
            else:
                raise CliError(
                    "unsupported_conversion", f"pdf converts to txt, png or jpg here, not {target}."
                )
        elif source == "csv" and target == "xlsx":
            outputs = [_csv_to_xlsx(src, out_dir, stem)]
        elif source == "xlsx" and target == "csv":
            outputs = [_xlsx_to_csv(src, out_dir, stem, args.sheet)]
        elif source in OFFICE_INPUTS:
            if source == "docx" and target == "md":
                target_path = _unique(out_dir / f"{stem}.md")
                target_path.write_text(_docx_to_md(src), encoding="utf-8")
                outputs = [_output(target_path, "md")]
            elif target in IMAGE_FORMATS:
                pdf = _office(src, "pdf", workdir)
                paths = _pdf_to_images(pdf, out_dir, stem, target)
                outputs = [_output(Path(p), target) for p in paths]
            elif target in (
                "pdf",
                "docx",
                "odt",
                "rtf",
                "html",
                "txt",
                "xlsx",
                "pptx",
                "csv",
                "odp",
                "ods",
            ):
                if target == "pdf":
                    outputs = [_finish_pdf(src, out_dir, stem, workdir)]
                else:
                    produced = _office(src, target, workdir)
                    final = _unique(out_dir / f"{stem}.{target}")
                    shutil.move(str(produced), final)
                    outputs = [_output(final, target)]
            else:
                raise CliError(
                    "unsupported_conversion", f"{source} cannot convert to {target} here."
                )
        else:
            raise CliError("unsupported_conversion", f"Cannot convert {source} to {target}.")
    finally:
        shutil.rmtree(workdir, ignore_errors=True)
    return {"ok": True, "action": "convert", "from": source, "to": target, "outputs": outputs}


def _output(path: Path, fmt: str, image: Any = None) -> dict[str, Any]:
    entry: dict[str, Any] = {"path": str(path), "format": fmt, "bytes": path.stat().st_size}
    if image is not None:
        entry["width"], entry["height"] = image.size
    if fmt == "pdf":
        entry["pages"] = _pdf_pages(path)
    return entry


def _write_docx_output(
    spec: dict[str, Any], out_dir: Path, stem: str, workdir: Path
) -> dict[str, Any]:
    target = _unique(out_dir / f"{stem}.docx")
    stats = _build_docx(spec, target, workdir)
    entry = _output(target, "docx")
    entry["stats"] = stats
    return entry


def _finish_pdf(src: Path, out_dir: Path, stem: str, workdir: Path) -> dict[str, Any]:
    produced = _office(src, "pdf", workdir)
    final = _unique(out_dir / f"{stem}.pdf")
    shutil.move(str(produced), final)
    return _output(final, "pdf")


def _csv_to_xlsx(src: Path, out_dir: Path, stem: str) -> dict[str, Any]:
    openpyxl = _module("openpyxl", "openpyxl")
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    with src.open(newline="", encoding="utf-8-sig") as handle:
        for row in csv.reader(handle):
            converted: list[Any] = []
            for cell in row:
                try:
                    converted.append(
                        int(cell)
                        if re.fullmatch(r"-?\d+", cell)
                        else float(cell)
                        if re.fullmatch(r"-?\d+\.\d+", cell)
                        else cell
                    )
                except ValueError:
                    converted.append(cell)
            sheet.append(converted)
    for cell in sheet[1] if sheet.max_row else []:
        cell.font = openpyxl.styles.Font(bold=True)
    target = _unique(out_dir / f"{stem}.xlsx")
    workbook.save(target)
    return _output(target, "xlsx")


def _xlsx_to_csv(src: Path, out_dir: Path, stem: str, sheet_name: str | None) -> dict[str, Any]:
    openpyxl = _module("openpyxl", "openpyxl")
    workbook = openpyxl.load_workbook(src, data_only=True, read_only=True)
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            raise CliError(
                "unknown_sheet", f"No sheet named {sheet_name!r}. Sheets: {workbook.sheetnames}."
            )
        sheet = workbook[sheet_name]
    else:
        sheet = workbook.worksheets[0]
    target = _unique(out_dir / f"{stem}.csv")
    with target.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        for row in sheet.iter_rows(values_only=True):
            writer.writerow(["" if v is None else v for v in row])
    workbook.close()
    return _output(target, "csv")


def do_create(args: argparse.Namespace) -> dict[str, Any]:
    spec = _load_spec(args.spec)
    if not isinstance(spec, dict):
        raise CliError("bad_spec", "create needs a JSON object with a 'blocks' list.")
    formats = {
        f.strip().lower().lstrip(".") for f in str(args.to or "docx").split(",") if f.strip()
    }
    unknown = formats - {"docx", "pdf"}
    if unknown:
        raise CliError("unsupported_format", f"create writes docx and pdf, not {sorted(unknown)}.")
    out_dir = _out_dir(args.out_dir)
    name = _safe_name(args.name or spec.get("name") or spec.get("title"), "document")
    workdir = Path(tempfile.mkdtemp(prefix="doccreate-"))
    try:
        outputs = [_write_docx_output(spec, out_dir, name, workdir)]
        if "pdf" in formats:
            outputs.append(_finish_pdf(Path(outputs[0]["path"]), out_dir, name, workdir))
    finally:
        shutil.rmtree(workdir, ignore_errors=True)
    return {"ok": True, "action": "create", "outputs": outputs}


def do_image(args: argparse.Namespace) -> dict[str, Any]:
    op = str(args.op or "").lower()
    if op not in ("fetch", "crop", "resize", "convert"):
        raise CliError("bad_op", "image needs op: fetch, crop, resize or convert.")
    if op == "fetch":
        if not args.url:
            raise CliError("missing_argument", "image op=fetch needs url.")
        source = args.url
        label = Path(args.url.split("?")[0]).stem or "image"
    else:
        if not args.file:
            raise CliError("missing_argument", f"image op={op} needs file.")
        if _URL.match(str(args.file)):
            source = args.file
            label = Path(args.file.split("?")[0]).stem or "image"
        else:
            source = str(_input_file(args.file))
            label = Path(source).stem
    image, original_fmt = _open_image(source)
    original_size = image.size
    if op == "crop":
        crop = _parse_crop(args.crop)
        if crop is None:
            raise CliError(
                "missing_argument", "image op=crop needs crop: 'x,y,w,h' or an aspect like 16:9."
            )
        image = _crop(image, crop)
    if op == "resize":
        if not args.width_px:
            raise CliError("missing_argument", "image op=resize needs width_px.")
        target_w = int(args.width_px)
        target_h = max(1, round(image.size[1] * target_w / image.size[0]))
        image = image.resize((target_w, target_h), _pil().Resampling.LANCZOS)
    if op == "convert" and not args.to:
        raise CliError("missing_argument", "image op=convert needs to, e.g. png, jpg or webp.")
    ext = (
        str(args.to or ("jpg" if original_fmt == "JPEG" else original_fmt.lower()))
        .lower()
        .lstrip(".")
    )
    if ext not in IMAGE_FORMATS:
        raise CliError("bad_format", f"Unsupported image format {ext!r}.")
    out_dir = _out_dir(args.out_dir)
    target = _unique(out_dir / f"{_safe_name(args.name, label + '-' + op)}.{ext}")
    _save_image(image, target, ext)
    return {
        "ok": True,
        "action": "image",
        "op": op,
        "path": str(target),
        "format": ext,
        "source_size": list(original_size),
        "width": image.size[0],
        "height": image.size[1],
    }


def do_chart(args: argparse.Namespace) -> dict[str, Any]:
    spec = _load_spec(args.spec)
    if not isinstance(spec, dict):
        raise CliError("bad_spec", "chart needs a JSON object with kind, labels and series.")
    out_dir = _out_dir(args.out_dir)
    target = _unique(out_dir / f"{_safe_name(args.name or spec.get('title'), 'chart')}.png")
    _render_chart(spec, target)
    image, _fmt = _open_image(str(target))
    return {
        "ok": True,
        "action": "chart",
        "path": str(target),
        "width": image.size[0],
        "height": image.size[1],
    }


def do_insert(args: argparse.Namespace) -> dict[str, Any]:
    from docx import Document

    docx_path = _input_file(args.file, "file (a .docx to insert into)")
    if docx_path.suffix.lower() != ".docx":
        raise CliError("unsupported_format", "insert adds pictures to .docx files.")
    if not args.image:
        raise CliError("missing_argument", "insert needs image: a path or URL.")
    position = str(args.position or "end")
    workdir = Path(tempfile.mkdtemp(prefix="docins-"))
    try:
        picture = _materialise(args.image, args.crop, workdir)
        document = Document(docx_path)
        before = len(document.inline_shapes)
        width_cm = min(float(args.width_cm or 12), _usable_cm(document))
        if position == "header":
            paragraph = document.sections[0].header.paragraphs[0]
            paragraph.add_run().add_picture(str(picture), width=_cm(width_cm))
        elif position == "end":
            paragraph = document.add_paragraph()
            paragraph.alignment = _align("center")
            paragraph.add_run().add_picture(str(picture), width=_cm(width_cm))
        elif position == "top":
            paragraph = document.add_paragraph()
            paragraph.alignment = _align("center")
            paragraph.add_run().add_picture(str(picture), width=_cm(width_cm))
            first = document.paragraphs[0]._p if document.paragraphs else None
            if first is not None and first is not paragraph._p:
                first.addprevious(paragraph._p)
        elif position.startswith("after:"):
            needle = position[len("after:") :].strip().lower()
            if not needle:
                raise CliError("bad_position", "position 'after:' needs the text to insert after.")
            anchor = next((p for p in document.paragraphs if needle in p.text.lower()), None)
            if anchor is None:
                raise CliError(
                    "anchor_not_found",
                    f"No paragraph contains {needle!r}.",
                    "Check the text against the document, or use position 'end' or 'header'.",
                )
            paragraph = document.add_paragraph()
            paragraph.alignment = _align("center")
            paragraph.add_run().add_picture(str(picture), width=_cm(width_cm))
            anchor._p.addnext(paragraph._p)
        else:
            raise CliError(
                "bad_position", "position must be 'top', 'end', 'header' or 'after:<text>'."
            )
        out_dir = _out_dir(args.out_dir)
        target = _unique(out_dir / f"{_safe_name(args.name, docx_path.stem + '-with-image')}.docx")
        document.save(target)
    finally:
        shutil.rmtree(workdir, ignore_errors=True)
    after = len(Document(target).inline_shapes)
    if after != before + 1:
        raise CliError(
            "verify_failed", f"Expected {before + 1} pictures after insert, found {after}."
        )
    return {
        "ok": True,
        "action": "insert",
        "path": str(target),
        "position": position,
        "width_cm": round(width_cm, 2),
        "pictures_before": before,
        "pictures_after": after,
    }


# --- Install and doctor --------------------------------------------------- #
def _tail(text: str, limit: int = 1500) -> str:
    return text.strip()[-limit:]


def _version_of(distribution: str) -> str | None:
    try:
        from importlib import metadata

        return metadata.version(distribution)
    except Exception:  # noqa: BLE001
        return None


def _doctor() -> dict[str, Any]:
    _ensure_site_on_path()
    modules: dict[str, str | None] = {}
    for import_name, distribution in MODULES:
        try:
            importlib.import_module(import_name)
            modules[distribution] = _version_of(distribution) or "installed"
        except Exception:  # noqa: BLE001
            modules[distribution] = None
    binaries = {
        name: shutil.which(name)
        for name in (
            "soffice",
            "libreoffice",
            "pdftoppm",
            "pdftotext",
            "pdfinfo",
            "convert",
            "google-chrome",
        )
    }
    core_ready = all(modules[d] for _i, d in MODULES)
    office_ready = bool(binaries["soffice"] or binaries["libreoffice"]) and bool(
        binaries["pdftoppm"]
    )
    return {
        "python": sys.version.split()[0],
        "cli_version": CLI_VERSION,
        "home": str(HOME),
        "modules": modules,
        "binaries": binaries,
        "core_ready": core_ready,
        "office_ready": office_ready,
        "next": None if core_ready else "Run document_forge action='install' with background=true.",
    }


def _run_step(
    command: list[str], timeout: int, env: dict[str, str] | None = None
) -> subprocess.CompletedProcess:
    return subprocess.run(command, capture_output=True, text=True, timeout=timeout, env=env)


def _pip_install() -> dict[str, Any]:
    SITE.mkdir(parents=True, exist_ok=True)
    base = [
        sys.executable,
        "-m",
        "pip",
        "install",
        "--quiet",
        "--disable-pip-version-check",
        "--upgrade",
        "--target",
        str(SITE),
        *PIP_PACKAGES,
    ]
    proc = _run_step(base, timeout=1500)
    output = (proc.stdout or "") + (proc.stderr or "")
    if proc.returncode != 0 and "externally-managed" in output:
        proc = _run_step(base + ["--break-system-packages"], timeout=1500)
        output = (proc.stdout or "") + (proc.stderr or "")
    if proc.returncode != 0 and (
        "No module named pip" in output or "No module named 'pip'" in output
    ):
        _run_step([sys.executable, "-m", "ensurepip", "--upgrade"], timeout=600)
        proc = _run_step(base, timeout=1500)
        output = (proc.stdout or "") + (proc.stderr or "")
    return {"ok": proc.returncode == 0, "tail": _tail(output, 800)}


def _apt_install() -> dict[str, Any]:
    if shutil.which("soffice") and shutil.which("pdftoppm"):
        return {"ok": True, "skipped": "LibreOffice and poppler already present"}
    if not shutil.which("apt-get"):
        return {"ok": False, "skipped": "apt-get is not available in this sandbox"}
    if os.geteuid() == 0:
        prefix: list[str] = []
    elif shutil.which("sudo") and _run_step(["sudo", "-n", "true"], timeout=20).returncode == 0:
        prefix = ["sudo", "-n"]
    else:
        return {"ok": False, "skipped": "not root and no passwordless sudo, so apt cannot run"}
    env = {**os.environ, "DEBIAN_FRONTEND": "noninteractive"}
    update = _run_step([*prefix, "apt-get", "update", "-qq"], timeout=900, env=env)
    install = _run_step(
        [*prefix, "apt-get", "install", "-y", "-qq", "--no-install-recommends", *APT_PACKAGES],
        timeout=2400,
        env=env,
    )
    return {
        "ok": install.returncode == 0,
        "update_ok": update.returncode == 0,
        "tail": _tail((install.stdout or "") + (install.stderr or ""), 800),
    }


def do_install() -> dict[str, Any]:
    pip = _pip_install()
    apt = _apt_install()
    doctor = _doctor()
    ok = bool(doctor["core_ready"]) and pip["ok"]
    result = {"ok": ok, "action": "install", "pip": pip, "apt": apt, "doctor": doctor}
    if not ok:
        raise CliError(
            "install_incomplete", "Some Python packages did not install.", pip.get("tail")
        )
    return result


def do_job(args: argparse.Namespace) -> dict[str, Any]:
    if not args.job_id:
        JOBS.mkdir(parents=True, exist_ok=True)
        records = []
        for path in sorted(JOBS.glob("*.json"), key=lambda p: p.stat().st_mtime, reverse=True)[:20]:
            try:
                record = json.loads(path.read_text())
            except (OSError, json.JSONDecodeError):
                continue
            records.append(
                {
                    "id": record.get("id"),
                    "action": record.get("action"),
                    "state": record.get("state"),
                }
            )
        return {"ok": True, "action": "job", "jobs": records}
    record = _read_job(args.job_id)
    if record is None:
        raise CliError(
            "unknown_job", f"No job {args.job_id}.", "Use action='job' without job_id to list jobs."
        )
    if record.get("state") == "running" and not _pid_alive(record.get("pid")):
        record["state"] = "lost"
        record["error"] = {
            "code": "process_ended",
            "message": "The job process ended without recording a result (the sandbox may have restarted).",
            "next": "Run the action again.",
        }
        _write_job(args.job_id, record)
    log = Path(record.get("log", ""))
    if record.get("state") != "done" and log.is_file():
        record["log_tail"] = _tail(log.read_text(errors="replace"), 1500)
    return {"ok": True, "action": "job", "job": record}


def _pid_alive(pid: Any) -> bool:
    if not isinstance(pid, int):
        return False
    try:
        os.kill(pid, 0)
    except OSError:
        return False
    return True


def _write_job(job_id: str, record: dict[str, Any]) -> None:
    JOBS.mkdir(parents=True, exist_ok=True)
    path = JOBS / f"{job_id}.json"
    tmp = path.with_suffix(".tmp")
    tmp.write_text(json.dumps(record, default=str, indent=2))
    os.replace(tmp, path)


def _read_job(job_id: str) -> dict[str, Any] | None:
    if not re.fullmatch(r"[0-9a-f]{12}", job_id or ""):
        raise CliError("bad_job_id", f"Job ids are 12 hex characters, got {job_id!r}.")
    path = JOBS / f"{job_id}.json"
    if not path.is_file():
        return None
    return json.loads(path.read_text())


def _launch_background(argv: list[str], action: str) -> dict[str, Any]:
    job_id = uuid.uuid4().hex[:12]
    JOBS.mkdir(parents=True, exist_ok=True)
    log = JOBS / f"{job_id}.log"
    _write_job(
        job_id,
        {"id": job_id, "action": action, "state": "queued", "started_at": _now(), "log": str(log)},
    )
    child = [sys.executable, os.path.abspath(__file__), *argv, "--job-id", job_id]
    with open(log, "ab") as handle:
        subprocess.Popen(
            child,
            stdin=subprocess.DEVNULL,
            stdout=handle,
            stderr=handle,
            start_new_session=True,
            close_fds=True,
        )
    return {
        "ok": True,
        "action": action,
        "job_id": job_id,
        "state": "running",
        "next": "Poll with document_forge action='job' and this job_id. The record turns "
        "to done (with the output paths) or failed (with the error).",
    }


def _dispatch(args: argparse.Namespace) -> dict[str, Any]:
    action = args.action
    handlers = {
        "doctor": lambda: {"ok": True, "action": "doctor", **_doctor()},
        "install": do_install,
        "job": lambda: do_job(args),
        "convert": lambda: do_convert(args),
        "create": lambda: do_create(args),
        "image": lambda: do_image(args),
        "chart": lambda: do_chart(args),
        "insert": lambda: do_insert(args),
    }
    handler = handlers.get(action)
    if handler is None:
        raise CliError("unknown_action", f"Unknown action {action!r}.")
    return handler()


def _run_child(args: argparse.Namespace) -> int:
    job_id = args.job_id
    record = _read_job(job_id) or {"id": job_id, "action": args.action}
    record.update(state="running", pid=os.getpid(), started_at=record.get("started_at") or _now())
    _write_job(job_id, record)
    code = 0
    try:
        record.update(state="done", result=_dispatch(args), error=None)
    except CliError as exc:
        record.update(
            state="failed", error={"code": exc.code, "message": exc.message, "next": exc.next_step}
        )
        code = 1
    except Exception as exc:  # noqa: BLE001 - recorded so the poll shows it
        record.update(
            state="failed",
            error={
                "code": "unexpected",
                "message": f"{type(exc).__name__}: {exc}",
                "traceback": _tail(traceback.format_exc(), 1200),
            },
        )
        code = 1
    record["finished_at"] = _now()
    _write_job(job_id, record)
    _emit(record)
    return code


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(prog="document_forge_cli", description=__doc__.split("\n")[0])
    parser.add_argument(
        "action",
        choices=[
            "doctor",
            "install",
            "job",
            "convert",
            "create",
            "image",
            "chart",
            "insert",
        ],
    )
    parser.add_argument("--spec", help="JSON object or path to a .json file (create, chart).")
    parser.add_argument("--file", help="Input file (convert, image, insert).")
    parser.add_argument(
        "--to", help="Target format (convert: pdf, docx, png ...; create: docx,pdf)."
    )
    parser.add_argument("--url", help="http(s) URL to fetch (image op=fetch).")
    parser.add_argument("--image", help="Picture to insert: a path or URL (insert).")
    parser.add_argument("--op", help="image operation: fetch, crop, resize or convert.")
    parser.add_argument(
        "--crop", help="'x,y,w,h' pixels or an aspect such as 16:9 (image, insert)."
    )
    parser.add_argument(
        "--width-px", dest="width_px", type=int, help="Target width in pixels (image op=resize)."
    )
    parser.add_argument(
        "--width-cm", dest="width_cm", type=float, help="Picture width in cm (insert)."
    )
    parser.add_argument("--position", help="insert position: top, end, header or after:<text>.")
    parser.add_argument("--name", help="Base name for the output file.")
    parser.add_argument(
        "--out-dir", dest="out_dir", help="Where to write. Defaults to ~/document_outputs."
    )
    parser.add_argument("--sheet", help="Sheet name for xlsx to csv.")
    parser.add_argument(
        "--background", action="store_true", help="Run detached and return a job id."
    )
    parser.add_argument(
        "--job-id", dest="job_id", help="Internal: the job this process records into."
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    raw = list(sys.argv[1:] if argv is None else argv)
    args = build_parser().parse_args(raw)
    _ensure_site_on_path()
    if args.job_id and args.action != "job":
        return _run_child(args)
    if args.background and args.action not in ("doctor", "job"):
        clean = [item for item in raw if item != "--background"]
        try:
            _emit(_launch_background(clean, args.action))
        except OSError as exc:
            _emit(
                {"ok": False, "action": args.action, "error": "launch_failed", "message": str(exc)}
            )
            return 1
        return 0
    try:
        result = _dispatch(args)
    except CliError as exc:
        _emit(
            {
                "ok": False,
                "action": args.action,
                "error": exc.code,
                "message": exc.message,
                "next": exc.next_step,
            }
        )
        return 1
    _emit(result)
    return 0


if __name__ == "__main__":
    sys.exit(main())
